"""Bounded upload ingestion and immutable SQLite upload snapshots."""
import csv
import io
import itertools
import math
import sqlite3
import zipfile
from contextlib import closing

from .config import MiB

ALIASES = {
    'compound': ('compound', 'compound_id', 'compound_name', 'drug', 'drug_id', 'molecule', 'molecule_id', 'ligand', 'id'),
    'target': ('target', 'target_id', 'target_name', 'protein', 'protein_id', 'gene', 'gene_symbol', 'targets', 'target_names', 'uniprot', 'uniprot_id', 'uniprot_accession', 'accession', 'uniprot_acc', 'protein_accession', 'target_accession'),
    'affinity': ('affinity', 'affinity_pkd', 'affinity_value', 'pkd', 'pic50', 'pki', 'value', 'score', 'activity', 'potency'),
    'price': ('price', 'price_usd_per_mg', 'price_usd_mg', 'price_per_mg', 'usd_per_mg', 'price_1mg', 'cost', 'value'),
}


def upload_rows(path, name, policy):
    """Yield actual rows, never trusting worksheet dimension declarations."""
    suffix = name.lower().rsplit('.', 1)[-1]
    if suffix == 'csv':
        csv.field_size_limit(16384)
        with open(path, encoding='utf-8-sig', newline='') as source:
            yield from csv.reader(source)
    elif suffix == 'xlsx':
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > 10000 or sum(item.file_size for item in entries) > policy.xlsx_bytes:
                raise ValueError('Excel archive exceeds the decompressed size limit.')
            # Fully consume the bounded archive to verify declared sizes/CRC.
            total = 0
            for item in entries:
                with archive.open(item) as stream:
                    while chunk := stream.read(MiB):
                        total += len(chunk)
                        if total > policy.xlsx_bytes:
                            raise ValueError('Excel archive exceeds the decompressed size limit.')
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        try:
            sheet = book.worksheets[0]
            sheet.reset_dimensions()
            yield from sheet.iter_rows(values_only=True)
        finally:
            book.close()
    else:
        raise ValueError('Use CSV or Excel (.xlsx). Legacy .xls is not supported.')


def text(value):
    value = str(value).strip() if value is not None else ''
    if len(value) > 4096 or '\x00' in value:
        raise ValueError('Identifiers must contain at most 4096 characters and no null bytes.')
    return value


def initialize(path, previous=None):
    if previous:
        with closing(sqlite3.connect(previous.as_uri() + '?mode=ro', uri=True)) as source, closing(sqlite3.connect(path)) as dest:
            source.backup(dest)
    with closing(sqlite3.connect(path)) as db:
        db.executescript('''
            PRAGMA journal_mode=DELETE;
            CREATE TABLE IF NOT EXISTS files(kind TEXT,name TEXT,size INTEGER,PRIMARY KEY(kind,name));
            CREATE TABLE IF NOT EXISTS affinity(file TEXT,compound TEXT,target TEXT,value REAL,UNIQUE(file,compound,target,value));
            CREATE TABLE IF NOT EXISTS prices(file TEXT,compound TEXT,value REAL,UNIQUE(file,compound));
            CREATE INDEX IF NOT EXISTS price_compound ON prices(compound);
            CREATE TABLE IF NOT EXISTS targets(file TEXT,target TEXT,UNIQUE(file,target));
            CREATE INDEX IF NOT EXISTS affinity_compound ON affinity(compound);
            CREATE INDEX IF NOT EXISTS affinity_target ON affinity(target);
        ''')


def ingest(path, staged, kind, policy, checkpoint=lambda: None):
    table = {'affinity': 'affinity', 'prices': 'prices', 'targets': 'targets'}[kind]
    with closing(sqlite3.connect(path)) as db, db:
        for item in staged:
            name = item['name']
            db.execute(f'DELETE FROM {table} WHERE file=?', (name,))
            db.execute('INSERT INTO files VALUES (?,?,?) ON CONFLICT(kind,name) DO UPDATE SET size=excluded.size', (kind, name, item['size']))
            if db.execute('SELECT count(*) FROM files').fetchone()[0] > policy.files:
                raise ValueError('Too many retained upload files.')
            if db.execute('SELECT coalesce(sum(size),0) FROM files').fetchone()[0] > policy.upload_bytes:
                raise ValueError('The session upload byte limit was exceeded.')
            rows = iter(upload_rows(item['path'], name, policy))
            header = next(rows, None)
            if not header:
                raise ValueError('The uploaded table is empty.')
            if len(header) > 32:
                raise ValueError('Upload tables may contain at most 32 columns.')
            header = [str(value or '').strip().lower().replace(' ', '_').replace('-', '_') for value in header]
            keys = {'affinity': ('compound','target','affinity'), 'prices': ('compound','price'), 'targets': ('target',)}[kind]
            indices = []
            for fallback, key in enumerate(keys):
                index = next((i for i, col in enumerate(header) if col in ALIASES[key]), None)
                if index is None:
                    if len(header) != len(keys):
                        raise ValueError(f'Missing {key} column in {name}.')
                    index = fallback
                indices.append(index)
            count = 0
            while batch := list(itertools.islice(rows, 1000)):
                checkpoint()
                values = []
                for row in batch:
                    if not any(v is not None and str(v).strip() for v in row):
                        continue
                    count += 1
                    if count > policy.upload_rows or len(row) > 32:
                        raise ValueError('Upload row or column limit exceeded.')
                    try:
                        selected = [row[i] for i in indices]
                    except IndexError:
                        raise ValueError('An uploaded row is missing required columns.') from None
                    if kind == 'targets':
                        target = text(selected[0])
                        if target:
                            values.append((name, target))
                    else:
                        ids = [text(v) for v in selected[:-1]]
                        if not all(ids):
                            continue
                        try:
                            number = float(selected[-1])
                        except (ValueError, TypeError):
                            raise ValueError('Affinity and price values must be finite numbers.') from None
                        if not math.isfinite(number) or (kind == 'prices' and number < 0):
                            raise ValueError('Values must be finite; prices must be non-negative.')
                        values.append((name, *ids, number))
                placeholders = ','.join('?' for _ in range(len(keys)+1))
                verb = 'INSERT OR REPLACE' if kind == 'prices' else 'INSERT OR IGNORE'
                db.executemany(f'{verb} INTO {table} VALUES ({placeholders})', values)
                # Reject during ingestion, before the next batch is allocated.
                total = db.execute(f'SELECT count(*) FROM {table}').fetchone()[0]
                ceiling = policy.upload_rows if kind == 'affinity' else (policy.compounds if kind == 'prices' else policy.targets)
                if total > ceiling:
                    raise ValueError(f'The session {kind} row limit ({ceiling:,}) was exceeded.')
            if not db.execute(f'SELECT 1 FROM {table} WHERE file=? LIMIT 1',(name,)).fetchone():
                raise ValueError(f'No valid data rows found in {name}.')
        if kind == 'affinity':
            n, t = db.execute('SELECT count(DISTINCT compound),count(DISTINCT target) FROM affinity').fetchone()
            policy.dimensions(n, t)


def edit(path, kind, operation, value):
    table = {'affinity':'affinity','prices':'prices','targets':'targets'}[kind]
    with closing(sqlite3.connect(path)) as db, db:
        if operation == 'clear':
            db.execute(f'DELETE FROM {table}')
            db.execute('DELETE FROM files WHERE kind=?', (kind,))
        elif operation == 'file':
            db.execute(f'DELETE FROM {table} WHERE file=?', (value,))
            db.execute('DELETE FROM files WHERE kind=? AND name=?', (kind,value))
        elif operation in ('compound','target'):
            db.execute(f'DELETE FROM {table} WHERE {operation}=?', (value,))
        else:
            raise ValueError('Unknown upload edit.')


def summary(path, kind, offset=0, limit=100):
    with closing(sqlite3.connect(path.as_uri()+'?mode=ro', uri=True)) as db:
        if kind == 'affinity':
            n,t,count = db.execute('SELECT count(DISTINCT compound),count(DISTINCT target),count(*) FROM (SELECT DISTINCT compound,target,value FROM affinity)').fetchone()
            compounds = [r[0] for r in db.execute('SELECT DISTINCT compound FROM affinity ORDER BY compound LIMIT ? OFFSET ?', (limit,offset))]
            targets = [r[0] for r in db.execute('SELECT DISTINCT target FROM affinity ORDER BY target')]
            files = [{'name':name,'num_datapoints':rows,'num_compounds':compounds_n,'num_targets':targets_n,'compounds':[],'targets':[]} for name,rows,compounds_n,targets_n in db.execute('SELECT file,count(*),count(DISTINCT compound),count(DISTINCT target) FROM affinity GROUP BY file')]
            return {'num_compounds':n,'num_targets':t,'num_datapoints':count,'compounds':compounds,'targets':targets,'all_files':files,'offset':offset,'limit':limit,'total':n,'paged':True}
        if kind == 'prices':
            total = db.execute('SELECT count(DISTINCT compound) FROM prices').fetchone()[0]
            compounds = [r[0] for r in db.execute('SELECT DISTINCT compound FROM prices ORDER BY compound LIMIT ? OFFSET ?', (limit,offset))]
            files = [{'name':name,'num_prices':n,'compounds':[]} for name,n in db.execute('SELECT file,count(*) FROM prices GROUP BY file')]
            return {'num_prices':total,'compounds':compounds,'all_files':files,'filename':', '.join(f['name'] for f in files),'offset':offset,'limit':limit,'total':total,'paged':True}
        return [r[0] for r in db.execute('SELECT DISTINCT target FROM targets ORDER BY rowid')]
