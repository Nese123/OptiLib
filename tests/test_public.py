"""Public API, isolation and failure regressions without production databases."""
import io
import json
import os
import re
import sqlite3
import tempfile
import time
import unittest
from pathlib import Path
from unittest.mock import patch

import numpy as np
from openpyxl import load_workbook

from webapp.public.config import Policy, configure_security, GiB
from webapp.public.runtime import Runtime, Rejected
from webapp.public.server import create_app
from webapp.public.tasks import execute
from webapp.public.worker import Supervisor


class PublicTests(unittest.TestCase):
    def setUp(self):
        environment=patch.dict(os.environ,{'OPTILIB_ENV':'development','SESSION_COOKIE_SECURE':'false','TRUSTED_HOSTS':'','OPTILIB_ACCEL_REDIRECT':'false'})
        environment.start();self.addCleanup(environment.stop)
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup)
        self.policy=Policy(self.tmp.name)
        self.policy.free_bytes=1
        self.policy.starts=100
        self.app=create_app(self.policy)
        self.app.config.update(TESTING=True,WTF_CSRF_ENABLED=False)
        self.app.extensions['public_limiter'].enabled=False
        self.runtime=self.app.extensions['runtime']
        self.runtime.heartbeat(True)
        self.client=self.app.test_client()
        self.client.get('/optimize')
        with self.client.session_transaction() as session:
            self.sid=session['sid']

    def perform(self,response):
        self.assertEqual(response.status_code,202,response.json)
        job_id=response.json['job_id']
        with self.runtime.connect() as db:
            db.execute("UPDATE jobs SET status='running',started=? WHERE id=?",(time.time(),job_id))
        # Resolve unknown fixture identifiers without touching local databases.
        from webapp.core.resolution import resolve_compounds,resolve_targets
        with patch.object(self.policy,'chembl',Path(self.tmp.name)/'missing.db'), \
                patch('webapp.core.pricing._lookup_molport_prices',return_value=({},{})):
            execute(self.runtime,self.runtime.job(job_id))
        result=self.client.get('/api/jobs/'+job_id)
        self.assertEqual(result.json['status'],'complete',result.json)
        return result.json['result']

    def upload(self,kind,text):
        return self.perform(self.client.post('/api/upload-'+kind,data={'files[]':(io.BytesIO(text.encode()),kind+'.csv')}))

    def dataset(self):
        self.upload('affinity','Compound,Target,Affinity\nA,T1,9\nA,T2,1\nA,T3,2\nB,T1,1\nB,T2,9\nB,T3,2\nC,T1,1\nC,T2,2\nC,T3,9\n')
        self.upload('prices','Compound,Price\nA,10\nB,20\nC,30\n')
        self.perform(self.client.post('/api/build-matrix-from-affinity',json={}))

    def test_end_to_end_artifacts_tiles_and_formula_safe_exports(self):
        self.dataset()
        self.assertEqual(self.client.get('/api/dataset-info').json['num_drugs'],3)
        self.perform(self.client.post('/api/run',json={'pop_size':5,'max_gen':10,'allowed_miss_pct':0}))
        response=self.client.get('/api/results')
        self.assertEqual(response.status_code,200,response.json)
        self.assertEqual(response.json['total'],3)
        self.assertEqual(response.json['comparison']['library']['total_cost'],60)
        tile=self.client.get('/api/heatmap-data?row_count=1&column_count=2').json
        self.assertEqual(len(tile['matrix']),1)
        self.assertEqual(len(tile['matrix'][0]),2)
        self.assertEqual(tile['total_rows'],3)
        self.assertEqual(self.client.get('/api/heatmap-data?row_count=101').status_code,400)
        response=self.client.get('/api/download/library')
        duplicate=self.client.get('/api/download/library')
        self.assertEqual(response.json['job_id'],duplicate.json['job_id'])
        self.perform(response)
        result=self.client.get('/api/download/library')
        self.assertEqual(result.status_code,200)
        book=load_workbook(io.BytesIO(result.data),read_only=True)
        self.assertEqual(len(list(book.active.values)),4)
        book.close();result.close()
        with patch.dict(os.environ,{'OPTILIB_ACCEL_REDIRECT':'true'}):
            accelerated=self.client.get('/api/download/library')
            self.assertTrue(accelerated.headers['X-Accel-Redirect'].startswith('/_optilib_artifacts/'+self.sid+'/'))
            self.assertEqual(accelerated.data,b'')
            self.assertEqual(accelerated.headers['Cache-Control'],'private, no-store')
        state=self.runtime.state(self.sid)
        self.assertIsInstance(state['dataset'],str)
        self.assertNotIn('matrix',state)
        old=state['selection']
        self.perform(self.client.post('/api/select-solution',json={'index':0}))
        self.assertNotEqual(old,self.runtime.state(self.sid)['selection'])

    def test_csrf_cookie_and_host_configuration(self):
        self.app.config['WTF_CSRF_ENABLED']=True
        self.assertEqual(self.client.post('/api/reset').status_code,400)
        page=self.client.get('/optimize')
        token=re.search(r'name="csrf-token" content="([^"]+)"',page.text).group(1)
        self.assertEqual(self.client.post('/api/reset',headers={'X-CSRFToken':token}).status_code,200)
        from flask import Flask
        with patch.dict(os.environ,{'OPTILIB_ENV':'production','SECRET_KEY':'placeholder','TRUSTED_HOSTS':'example.org','SESSION_COOKIE_SECURE':'true'}):
            with self.assertRaises(ValueError):configure_security(Flask('invalid'))
        with patch.dict(os.environ,{'OPTILIB_ENV':'production','SECRET_KEY':'a'*32,'TRUSTED_HOSTS':'example.org','SESSION_COOKIE_SECURE':'true'}):
            app=Flask('valid');configure_security(app)
            @app.get('/')
            def home():return 'ok'
            client=app.test_client()
            self.assertEqual(client.get('/',base_url='https://wrong.example').status_code,400)
            self.assertEqual(client.get('/',base_url='https://example.org').status_code,200)

    def test_cross_session_access_and_atomic_retired_publication(self):
        response=self.client.post('/api/upload-prices',data={'files[]':(io.BytesIO(b'Compound,Price\nA,1\n'),'a.csv')})
        job=response.json['job_id']
        other=self.app.test_client()
        self.assertEqual(other.get('/api/jobs/'+job).status_code,404)
        self.runtime.retire(self.sid)
        with self.runtime.connect() as db:db.execute("UPDATE jobs SET status='running' WHERE id=?",(job,))
        self.assertFalse(self.runtime.complete(job,{'dataset':'evil'},{'status':'complete'}))
        self.assertNotIn('dataset',self.runtime.state(self.sid))

    def test_multipart_file_larger_than_decoder_buffer_is_accepted(self):
        content=b'Compound,Price\n'+b'COMPOUND,10\n'*60000
        response=self.client.post('/api/upload-prices',data={'files[]':(io.BytesIO(content),'large.csv')})
        self.assertEqual(response.status_code,202,response.json)

    def test_atomic_upload_limit_failure(self):
        self.upload('affinity','Compound,Target,Affinity\nA,T1,2\n')
        previous=self.runtime.state(self.sid)['uploads']
        self.policy.upload_rows=1
        response=self.client.post('/api/upload-affinity',data={'files[]':(io.BytesIO(b'Compound,Target,Affinity\nB,T2,3\n'),'second.csv')})
        job=response.json['job_id']
        with self.runtime.connect() as db:db.execute("UPDATE jobs SET status='running' WHERE id=?",(job,))
        with self.assertRaises(ValueError):execute(self.runtime,self.runtime.job(job))
        self.assertEqual(self.runtime.state(self.sid)['uploads'],previous)

    def test_readiness_rejects_empty_databases_liveness_survives(self):
        self.policy.database=Path(self.tmp.name)
        self.policy.chembl=Path(self.tmp.name)/'chembl.db'
        self.policy.chembl.touch();(Path(self.tmp.name)/'molport.db').touch()
        self.assertEqual(self.client.get('/health').status_code,503)
        self.assertEqual(self.client.get('/live').status_code,200)

    def test_admission_is_exclusive_for_large_jobs_and_per_ip(self):
        # Use the actual current epoch for additional sessions.
        with self.runtime.connect() as db:epoch=db.execute("SELECT value FROM meta WHERE key='epoch'").fetchone()[0]
        other=self.runtime.session(None,epoch)['sid']
        a=self.runtime.admit(self.sid,'ip1','selection',{},reserve=1)
        with self.assertRaises(Rejected) as error:self.runtime.admit(other,'ip1','selection',{},reserve=1)
        self.assertEqual(error.exception.status,429)
        with self.assertRaises(Rejected):self.runtime.admit(other,'ip2','optimization',{},cells=100000000,reserve=1)
        b=self.runtime.admit(other,'ip2','selection',{},reserve=1)
        self.assertNotEqual(a,b)

    def test_cleanup_preserves_unowned_directories(self):
        unrelated=Path(self.tmp.name)/'sessions'/'notes';unrelated.mkdir();(unrelated/'keep').write_text('keep')
        self.runtime.retire(self.sid);self.runtime.cleanup()
        self.assertTrue((unrelated/'keep').exists())

    def test_storage_reservations_reject_without_mutating_state(self):
        before=self.runtime.state(self.sid)
        self.policy.session_bytes=1
        with self.assertRaises(Rejected) as error:self.runtime.admit(self.sid,'ip','optimization',{},reserve=2)
        self.assertEqual(error.exception.status,413)
        self.assertEqual(before,self.runtime.state(self.sid))


class UploadArchiveTests(unittest.TestCase):
    def test_expansion_and_legacy_rejections(self):
        import zipfile
        from webapp.public.ingestion import upload_rows
        with tempfile.TemporaryDirectory() as tmp:
            policy=Policy(tmp);policy.xlsx_bytes=100
            path=Path(tmp)/'bomb.xlsx'
            with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as archive:
                archive.writestr('xl/worksheets/sheet1.xml','x'*1000)
            with self.assertRaisesRegex(ValueError,'decompressed'):
                list(upload_rows(path,'bomb.xlsx',policy))
            with self.assertRaisesRegex(ValueError,'Legacy'):
                list(upload_rows(path,'legacy.xls',policy))

    def test_formula_and_url_strings_are_exported_as_text(self):
        from webapp.core.storage import write_rows_excel
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'test.xlsx'
            write_rows_excel(['name'],[['=1+1'],['https://example.org']],path)
            book=load_workbook(path)
            self.assertEqual(book.active['A2'].value,'=1+1')
            self.assertEqual(book.active['A2'].data_type,'s')
            self.assertIsNone(book.active['A3'].hyperlink)
            book.close()


class SupervisorTests(unittest.TestCase):
    setUp = PublicTests.setUp
    def test_spawned_upload_process_completes_and_releases_capacity(self):
        response=self.client.post('/api/upload-prices',data={'files[]':(io.BytesIO(b'Compound,Price\nA,10\n'),'p.csv')})
        job=response.json['job_id']
        supervisor=Supervisor(self.runtime)
        self.addCleanup(supervisor.close)
        until=time.monotonic()+30
        while time.monotonic()<until:
            supervisor.tick()
            if self.runtime.job(job)['status'] not in ('reserved','running','stopping'):
                break
            time.sleep(.1)
        self.assertEqual(self.runtime.job(job)['status'],'complete',self.runtime.job(job))
        self.assertEqual(self.runtime.job(job)['reserved_bytes'],0)

    def test_deadline_requests_stop_then_forces_exit(self):
        from unittest.mock import Mock
        job=self.runtime.admit(self.sid,'ip','optimization',{},reserve=1)
        with self.runtime.connect() as db:
            db.execute("UPDATE jobs SET status='running',started=? WHERE id=?",(time.time()-15000,job))
        supervisor=Supervisor(self.runtime)
        process=Mock();process.is_alive.return_value=True;process.pid=987654321
        supervisor.children[job]=process
        with patch('webapp.public.worker.resident_bytes',return_value=0):
            supervisor.tick()
        self.assertEqual(self.runtime.job(job)['status'],'stopping')
        with self.runtime.connect() as db:
            db.execute('UPDATE jobs SET stop_at=? WHERE id=?',(time.time()-100,job))
        with patch('webapp.public.worker.os.killpg') as kill:
            supervisor.tick()
            kill.assert_called_once()
        self.assertEqual(self.runtime.job(job)['status'],'cancelled')
        self.assertEqual(self.runtime.job(job)['reserved_bytes'],0)

    def test_worker_restart_cancels_orphaned_running_jobs(self):
        job=self.runtime.admit(self.sid,'ip','selection',{},reserve=1)
        with self.runtime.connect() as db:
            db.execute("UPDATE jobs SET status='running',started=? WHERE id=?",(time.time(),job))
        Supervisor(self.runtime).tick()
        self.assertEqual(self.runtime.job(job)['status'],'cancelled')


class PublicationAndPolicyTests(unittest.TestCase):
    setUp = PublicTests.setUp
    perform = PublicTests.perform
    upload = PublicTests.upload
    dataset = PublicTests.dataset
    def test_concurrent_export_admission_reuses_one_reservation(self):
        from concurrent.futures import ThreadPoolExecutor
        import threading
        ready=threading.Barrier(2)
        def admit(_):
            ready.wait(timeout=5)
            return self.runtime.admit(self.sid,'ip','export',{'key':'matrix:csv:revision'},reserve=1)
        with ThreadPoolExecutor(max_workers=2) as executor:
            jobs=list(executor.map(admit,range(2)))
        self.assertEqual(jobs[0],jobs[1])
        with self.runtime.connect() as db:
            self.assertEqual(db.execute("SELECT count(*) FROM jobs WHERE status='reserved'").fetchone()[0],1)

    def test_reset_does_not_erase_hourly_ip_budget(self):
        self.policy.starts=1
        job=self.runtime.admit(self.sid,'limited-ip','optimization',{},reserve=1)
        self.runtime.fail(job)
        self.runtime.retire(self.sid);self.runtime.cleanup()
        with self.runtime.connect() as db:epoch=db.execute("SELECT value FROM meta WHERE key='epoch'").fetchone()[0]
        sid=self.runtime.session(None,epoch)['sid']
        with self.assertRaises(Rejected) as error:self.runtime.admit(sid,'limited-ip','optimization',{},reserve=1)
        self.assertEqual(error.exception.status,429)

    def test_stop_publishes_feasible_partial_results(self):
        self.dataset()
        response=self.client.post('/api/run',json={'pop_size':5,'max_gen':10,'allowed_miss_pct':0})
        job=response.json['job_id']
        with self.runtime.connect() as db:db.execute("UPDATE jobs SET status='running',started=? WHERE id=?",(time.time(),job))
        from webapp.public.tasks import Progress
        original=Progress.notify
        def stop_on_generation(callback,algorithm):
            if algorithm.n_gen>=2:
                self.runtime.stop(self.sid)
                callback.context.last_check=0
            original(callback,algorithm)
        with patch.object(Progress,'notify',stop_on_generation):
            execute(self.runtime,self.runtime.job(job))
        result=self.runtime.job(job)
        self.assertEqual(result['status'],'complete')
        self.assertTrue(result['result']['partial'])
        self.assertEqual(self.client.get('/api/results').status_code,200)

    def test_new_download_preparation_requires_csrf(self):
        self.dataset()
        self.app.config['WTF_CSRF_ENABLED']=True
        self.assertEqual(self.client.get('/api/download/matrix').status_code,400)

    def test_failed_job_preserves_existing_dataset(self):
        self.dataset()
        old=self.runtime.state(self.sid)['dataset']
        response=self.client.post('/api/build-matrix-from-affinity',json={'selectivity_threshold':1e10})
        job=response.json['job_id']
        with self.runtime.connect() as db:db.execute("UPDATE jobs SET status='running' WHERE id=?",(job,))
        with self.assertRaises(ValueError):execute(self.runtime,self.runtime.job(job))
        self.runtime.fail(job)
        self.assertEqual(self.runtime.state(self.sid)['dataset'],old)
        self.assertTrue(self.client.get('/api/dataset-info').json['ready'])


if __name__ == '__main__':
    unittest.main()
